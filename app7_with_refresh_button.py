import os
import csv
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import customtkinter
import tkinter as tk
from tkinter import END, messagebox
from tkinter.filedialog import askopenfilenames
from docx import Document
from PyPDF2 import PdfReader
from fpdf import FPDF
from difflib import SequenceMatcher
from PIL import Image, ImageDraw, ImageFont, ImageTk

# Create necessary folders
if not os.path.exists("Pending"):
    os.makedirs("Pending")
if not os.path.exists("Reports"):
    os.makedirs("Reports")
if not os.path.exists("Screenshots"):
    os.makedirs("Screenshots")

# File conversion functions (unchanged)
def convert_docx_to_txt(docx_file, txt_file):
    """Convert .docx to .txt using python-docx."""
    document = Document(docx_file)
    with open(txt_file, "w", encoding="utf-8") as txt:
        for paragraph in document.paragraphs:
            txt.write(paragraph.text + "\n")

def convert_pdf_to_txt(pdf_file, txt_file):
    """Convert .pdf to .txt using PyPDF2."""
    reader = PdfReader(pdf_file)
    with open(txt_file, "w", encoding="utf-8") as txt:
        for page in reader.pages:
            txt.write(page.extract_text() + "\n")

def convert_to_txt(input_file):
    """Determine file type and convert to .txt, saving in the Pending folder."""
    base_name = os.path.basename(os.path.splitext(input_file)[0])  # Get the file name without extension
    output_file = os.path.join("Pending", f"{base_name}.txt")
    ext = os.path.splitext(input_file)[-1].lower()
    try:
        if ext == ".docx":
            convert_docx_to_txt(input_file, output_file)
        elif ext == ".pdf":
            convert_pdf_to_txt(input_file, output_file)
        elif ext == ".txt":
            with open(input_file, "r", encoding="utf-8") as f_in:
                with open(output_file, "w", encoding="utf-8") as f_out:
                    f_out.write(f_in.read())
        else:
            raise ValueError(f"Unsupported file type: {ext}")
        print(f"Conversion successful: {input_file} -> {output_file}")
    except Exception as e:
        print(f"Error converting file: {e}")

# Global list to store paths of uploaded files
uploaded_files = []
plagiarism_results = []

# Function to upload multiple files (unchanged)
def upload_files():
    """Upload multiple files and convert them to .txt in the Pending folder."""
    filepaths = askopenfilenames(
        filetypes=[
            ("All files", "*.*")
        ]
    )
    if not filepaths:
        return

    global uploaded_files
    uploaded_files = []

    for filepath in filepaths:
        base_name = os.path.basename(filepath)
        txt_path = os.path.join("Pending", f"{os.path.splitext(base_name)[0]}.txt")
        convert_to_txt(filepath)
        uploaded_files.append(txt_path)

    messagebox.showinfo("Upload Complete", "Files uploaded and converted successfully.")

# Function to extract copied texts and their source documents (unchanged)
def extract_copied_texts(file1, file2, threshold=0.7):
    """Extract copied texts between two files and identify their source documents."""
    with open(file1, encoding='utf-8') as f1, open(file2, encoding='utf-8') as f2:
        text1 = f1.read()
        text2 = f2.read()

    matcher = SequenceMatcher(None, text1, text2)
    copied_texts = []

    for match in matcher.get_matching_blocks():
        if match.size >= threshold * len(text1):  # Only consider significant matches
            copied_text = text1[match.a:match.a + match.size].strip()
            if copied_text:  # Ignore empty matches
                copied_texts.append({
                    "source_document_1": os.path.basename(file1),
                    "source_document_2": os.path.basename(file2),
                    "copied_text": copied_text
                })

    return copied_texts

# Function to perform plagiarism check (unchanged)
def check_uploaded_files_plagiarism():
    """Perform plagiarism check on the uploaded .txt files in the Pending folder."""
    if not uploaded_files:
        messagebox.showwarning("No Files", "No files have been uploaded.")
        return

    try:
        # Create the "corrupt documents" folder if it doesn't exist
        corrupt_folder = os.path.join("Pending", "corrupt documents")
        if not os.path.exists(corrupt_folder):
            os.makedirs(corrupt_folder)

        # Read the content of the uploaded files
        notes = []
        valid_files = []
        corrupted_files = []

        for file in uploaded_files:
            try:
                with open(file, encoding='utf-8') as f:
                    content = f.read().strip()
                    if not content:
                        # Move the empty file to the "corrupt documents" folder
                        file_name = os.path.basename(file)
                        corrupt_file_path = os.path.join(corrupt_folder, file_name)
                        os.rename(file, corrupt_file_path)
                        corrupted_files.append(file_name)
                        messagebox.showwarning("Empty File", f"The file {file_name} is empty and has been moved to 'corrupt documents'.")
                        continue  # Skip this file and continue with the next one
                    notes.append(content)
                    valid_files.append(file)
            except Exception as e:
                # Handle file access errors (e.g., file is open in another program)
                file_name = os.path.basename(file)
                corrupt_file_path = os.path.join(corrupt_folder, file_name)
                os.rename(file, corrupt_file_path)
                corrupted_files.append(file_name)
                messagebox.showwarning("File Access Error", f"The file {file_name} could not be accessed and has been moved to 'corrupt documents'. Error: {e}")

        # Check if all files are empty or corrupted
        if not notes:
            messagebox.showerror("Empty Files", "All uploaded files are empty, corrupted, or contain no meaningful content.")
            return

        # Ask the user if they want to continue with the remaining files
        if len(valid_files) < len(uploaded_files):
            continue_check = messagebox.askyesno("Continue?", "Some files were corrupted or empty and moved to 'corrupt documents'. Do you want to continue with the remaining files?")
            if not continue_check:
                return

        # Perform TF-IDF vectorization and cosine similarity calculation
        vectors = TfidfVectorizer().fit_transform(notes).toarray()
        s_vectors = list(zip(valid_files, vectors))
        global plagiarism_results
        plagiarism_results = []

        for student_a, text_vector_a in s_vectors:
            new_vectors = s_vectors.copy()
            # Find the index of the current student_a in new_vectors using the file path (student_a)
            current_index = next(i for i, (file_path, _) in enumerate(new_vectors) if file_path == student_a)
            del new_vectors[current_index]
            for student_b, text_vector_b in new_vectors:
                sim_score = cosine_similarity([text_vector_a, text_vector_b])[0][1]
                student_pair = sorted((os.path.basename(student_a), os.path.basename(student_b)))
                plagiarism_results.append({
                    "Assignment 1": student_pair[0],
                    "Assignment 2": student_pair[1],
                    "Similarity Score": f"{sim_score:.2f}",
                    "Plagiarism Status": "Flagged" if sim_score >= 0.7 else "Clean"
                })

        # Notify the user about corrupted files and successful plagiarism check
        if corrupted_files:
            messagebox.showinfo(
                "Plagiarism Check Complete",
                f"Plagiarism check completed successfully.\n\n"
                f"The following files were corrupted or empty and moved to 'corrupt documents':\n"
                f"{', '.join(corrupted_files)}"
            )
        else:
            messagebox.showinfo("Plagiarism Check Complete", "Plagiarism check completed successfully.")

    except Exception as e:
        messagebox.showerror("Error", f"An error occurred during the plagiarism check: {e}")

# Function to display copied texts and their source documents (unchanged)
def show_copied_texts():
    """Display the copied texts and their source documents in a new window with an improved layout."""
    if not plagiarism_results:
        messagebox.showerror("No Results", "No plagiarism results to display.")
        return

    # Create a new window for displaying copied texts
    copied_texts_window = customtkinter.CTkToplevel()
    copied_texts_window.title("Copied Texts and Source Documents")
    copied_texts_window.geometry("1200x800")
    copied_texts_window.state("zoomed")  # Make the window full-screen

    # Create a main frame to hold the content
    main_frame = customtkinter.CTkFrame(copied_texts_window)
    main_frame.pack(fill="both", expand=True, padx=0, pady=0)  # Remove padding to fill the entire window

    # Create a canvas and a scrollbar for the main frame
    canvas = tk.Canvas(main_frame)
    scrollbar = customtkinter.CTkScrollbar(main_frame, orientation="vertical", command=canvas.yview)
    scrollable_frame = customtkinter.CTkFrame(canvas)

    # Configure the canvas to work with the scrollbar
    scrollable_frame.bind(
        "<Configure>",
        lambda e: canvas.configure(
            scrollregion=canvas.bbox("all")
        )
    )

    # Add the scrollable_frame to the canvas and set its width to match the canvas
    canvas.create_window((0, 0), window=scrollable_frame, anchor="nw", width=copied_texts_window.winfo_width())

    # Configure the canvas to use the scrollbar
    canvas.configure(yscrollcommand=scrollbar.set)

    # Bind mouse scroll event to the canvas
    def _on_mousewheel(event):
        canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

    canvas.bind_all("<MouseWheel>", _on_mousewheel)

    # Pack the canvas and scrollbar tightly together
    canvas.pack(side="left", fill="both", expand=True, padx=0, pady=0)  # Remove padding for canvas
    scrollbar.pack(side="right", fill="y", padx=0, pady=0)  # Remove padding for scrollbar

    # Extract and display copied texts for flagged documents
    flagged_documents = [result for result in plagiarism_results if result["Plagiarism Status"] == "Flagged"]
    if not flagged_documents:
        no_results_label = customtkinter.CTkLabel(scrollable_frame, text="No flagged documents found.", font=("Arial", 14))
        no_results_label.pack(pady=10)
        return

    for result in flagged_documents:
        file1 = os.path.join("Pending", result["Assignment 1"])
        file2 = os.path.join("Pending", result["Assignment 2"])

        # Extract copied texts
        copied_texts = extract_copied_texts(file1, file2)

        if copied_texts:
            # Create a frame for each flagged document pair
            document_frame = customtkinter.CTkFrame(scrollable_frame)
            document_frame.pack(fill="x", padx=10, pady=10)  # Keep padding for document frames

            # Add a heading for the source documents
            source_label = customtkinter.CTkLabel(
                document_frame,
                text=f"Source Documents: {result['Assignment 1']} and {result['Assignment 2']}",
                font=("Arial", 14, "bold"),
                text_color="#4CAF50" if customtkinter.get_appearance_mode() == "Dark" else "#2E2E2E"  # Green in dark mode, dark gray in light mode
            )
            source_label.pack(anchor="w", padx=10, pady=(10, 5))  # Keep padding for labels

            # Add a separator
            separator = tk.Frame(document_frame, height=2, bg="gray")
            separator.pack(fill="x", padx=10, pady=5)  # Keep padding for separators

            # Display each copied text
            for copied_text in copied_texts:
                # Create a frame for each copied text
                text_frame = customtkinter.CTkFrame(document_frame)
                text_frame.pack(fill="x", padx=20, pady=5)  # Keep padding for text frames

                # Add the copied text
                copied_text_label = customtkinter.CTkLabel(
                    text_frame,
                    text=f"Copied Text:\n{copied_text['copied_text']}",
                    font=("Arial", 12),
                    wraplength=1000,
                    justify="left",
                    text_color="#333333",  # Dark gray for better readability
                    fg_color="#F0F0F0",  # Light gray background for the text frame
                    corner_radius=5
                )
                copied_text_label.pack(anchor="w", padx=10, pady=5, fill="x")  # Keep padding for copied text

            # Add a separator between document pairs
            separator = tk.Frame(scrollable_frame, height=2, bg="gray")
            separator.pack(fill="x", padx=10, pady=10)  # Keep padding for separators

    # Update the canvas scroll region
    canvas.configure(scrollregion=canvas.bbox("all"))

# Function to export the plagiarism report (unchanged)
def export_report():
    """Generate and export plagiarism report to CSV and PDF."""
    if not plagiarism_results:
        messagebox.showerror("No Results", "No plagiarism results to export.")
        return

    # Export to CSV
    csv_filename = os.path.join("Reports", "plagiarism_report.csv")
    keys = plagiarism_results[0].keys()  # Assuming all reports have the same keys
    with open(csv_filename, mode='w', newline='', encoding='utf-8') as file:
        writer = csv.DictWriter(file, fieldnames=keys)
        writer.writeheader()
        writer.writerows(plagiarism_results)

    # Export to PDF
    pdf_filename = os.path.join("Reports", "plagiarism_report.pdf")
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Plagiarism Report", ln=True, align='C')
    pdf.ln(10)

    for result in plagiarism_results:
        pdf.multi_cell(0, 10, txt=f"{result}")
        pdf.ln(2)

    pdf.output(pdf_filename)
    messagebox.showinfo("Report Generated", f"Report saved successfully in Reports folder.\nCSV: {csv_filename}\nPDF: {pdf_filename}")

# Function to open the generated report (unchanged)
def open_report():
    """Allow the user to open the report in either CSV or PDF format."""
    report_choice = messagebox.askquestion("Open Report", "Would you like to open the report in PDF format? Click 'Yes' for PDF and 'No' for CSV.")
    if report_choice == 'yes':
        os.system(f'start Reports/plagiarism_report.pdf')
    else:
        os.system(f'start Reports/plagiarism_report.csv')

# Function to reset the application (unchanged)
def reset_application():
    """Clear all data and reset the application to its initial state."""
    global uploaded_files, plagiarism_results
    uploaded_files = []
    plagiarism_results = []

    # Clear Pending folder
    for file in os.listdir("Pending"):
        file_path = os.path.join("Pending", file)
        if os.path.isfile(file_path):
            os.remove(file_path)

    # Clear Screenshots folder
    for file in os.listdir("Screenshots"):
        file_path = os.path.join("Screenshots", file)
        if os.path.isfile(file_path):
            os.remove(file_path)

    messagebox.showinfo("Reset Complete", "Application has been reset to its initial state.")

# Function to toggle between dark and light modes
def toggle_appearance_mode():
    """Toggle between dark and light modes and update widget colors."""
    current_mode = customtkinter.get_appearance_mode()
    if current_mode == "Dark":
        customtkinter.set_appearance_mode("Light")
        btn_toggle_mode.configure(image=moon_icon)  # Show moon icon for light mode
    else:
        customtkinter.set_appearance_mode("Dark")
        btn_toggle_mode.configure(image=sun_icon)  # Show sun icon for dark mode
    update_widget_colors()

# Function to update widget colors based on the selected mode
def update_widget_colors():
    """Update the colors of all widgets based on the selected mode."""
    current_mode = customtkinter.get_appearance_mode()
    if current_mode == "Light":
        # Light mode colors
        window.configure(fg_color="#F5F5F5")  # Light gray background
        welcome_frm.configure(fg_color="#F5F5F5")  # Light gray background
        welcome_lbl.configure(text_color="#333333")  # Dark gray text
        button_frm.configure(fg_color="#F5F5F5")  # Light gray background

        # Buttons remain the same as in dark mode
        for button in [btn_upload, btn_check, btn_report, btn_open_report, btn_reset, btn_show_copied_texts]:
            button.configure(fg_color="#4CAF50", hover_color="#45a049", text_color="#FFFFFF")  # Green buttons with white text
        btn_reset.configure(fg_color="#FF0000", hover_color="#CC0000")  # Red reset button

        # Update toggle button background to match the app's background
        btn_toggle_mode.configure(fg_color="#F5F5F5", hover_color="#F5F5F5", text_color="#000000")
    else:
        # Dark mode colors (default)
        window.configure(fg_color="#2E2E2E")  # Dark gray background
        welcome_frm.configure(fg_color="#2E2E2E")  # Dark gray background
        welcome_lbl.configure(text_color="yellow")  # Yellow text
        button_frm.configure(fg_color="#2E2E2E")  # Dark gray background

        # Buttons remain the same as in dark mode
        for button in [btn_upload, btn_check, btn_report, btn_open_report, btn_reset, btn_show_copied_texts]:
            button.configure(fg_color="#4CAF50", hover_color="#45a049", text_color="#FFFFFF")  # Green buttons with white text
        btn_reset.configure(fg_color="#FF0000", hover_color="#CC0000")  # Red reset button

        # Update toggle button background to match the app's background
        btn_toggle_mode.configure(fg_color="#2E2E2E", hover_color="#2E2E2E", text_color="#FFFFFF")

# GUI IMPLEMENTED USING CUSTOMTKINTER
window = customtkinter.CTk()
customtkinter.set_appearance_mode("dark")  # Always open the app in dark mode
window.title("Plagiarism Checker System")

# Set the window to maximized state (covers the screen but leaves the taskbar visible)
window.state("zoomed")

# Load icons for the toggle button
sun_icon = ImageTk.PhotoImage(Image.open("sun.png").resize((32, 32)))  # Replace "sun.png" with the path to your sun icon
moon_icon = ImageTk.PhotoImage(Image.open("moon.png").resize((32, 32)))  # Replace "moon.png" with the path to your moon icon

# FRONT END DESIGN
# 1st Frame containing the welcome and intro msg label widgets.
welcome_frm = customtkinter.CTkFrame(window)
welcome_msg = "Welcome to the state of the art Plagiarism Checker System"
welcome_msg_variable = tk.StringVar(welcome_frm, welcome_msg)
welcome_lbl = customtkinter.CTkLabel(welcome_frm, textvariable=welcome_msg_variable,
                                     height=100, corner_radius=20, 
                                     text_color="yellow", font=("Comic Sans MS bold", 30))
welcome_lbl.grid(row=0, column=0, padx=200, pady=(20, 0), sticky="nsew")

# 3rd Frame contains the buttons to help with the functionality of the whole app.
button_frm = customtkinter.CTkFrame(window)
btn_upload = customtkinter.CTkButton(
    button_frm, corner_radius=30, hover_color="Green", text="UPLOAD FILES", command=upload_files
)
btn_upload.grid(row=0, column=0, padx=(180, 10), pady=10)

btn_check = customtkinter.CTkButton(
    button_frm, corner_radius=30, hover_color="Green", text="CHECK PLAGIARISM", command=check_uploaded_files_plagiarism
)
btn_check.grid(row=0, column=1, padx=10, pady=10)

btn_report = customtkinter.CTkButton(
    button_frm, corner_radius=30, hover_color="Green", text="EXPORT REPORT", command=export_report
)
btn_report.grid(row=0, column=2, padx=10, pady=10)

btn_open_report = customtkinter.CTkButton(
    button_frm, corner_radius=30, hover_color="Green", text="OPEN REPORT", command=open_report
)
btn_open_report.grid(row=0, column=3, padx=10, pady=10)

btn_reset = customtkinter.CTkButton(
    button_frm, corner_radius=30, hover_color="Red", text="RESET", command=reset_application
)
btn_reset.grid(row=0, column=4, padx=10, pady=10)

# Add the "SHOW COPIED TEXTS" button
btn_show_copied_texts = customtkinter.CTkButton(
    button_frm, corner_radius=30, hover_color="Green", text="SHOW COPIED TEXTS", command=show_copied_texts
)
btn_show_copied_texts.grid(row=0, column=5, padx=10, pady=10)

# Create a new frame for the top-right corner
top_right_frame = customtkinter.CTkFrame(window)
top_right_frame.grid(row=0, column=1, padx=20, pady=20, sticky="ne")  # Position in the top-right corner

# Add the "TOGGLE MODE" button to the top-right frame
btn_toggle_mode = customtkinter.CTkButton(
    top_right_frame, text="", image=sun_icon, command=toggle_appearance_mode,
    fg_color="#2E2E2E", hover_color="#2E2E2E", width=40, height=40  # Increase button size to 40x40
)
btn_toggle_mode.grid(row=0, column=0)

# Center the frames within the window
welcome_frm.grid(row=0, column=0, padx=20, pady=(50, 10), sticky="nsew")
button_frm.grid(row=1, column=0, padx=20, pady=(10, 50), sticky="nsew")

# Configure the grid to center the frames
window.grid_rowconfigure(0, weight=0)
window.grid_rowconfigure(0, weight=0)
window.grid_columnconfigure(0, weight=1)

# Initialize widget colors based on the default mode (dark)
update_widget_colors()

window.mainloop()